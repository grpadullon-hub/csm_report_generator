from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_TABLE_ALIGNMENT


# ============================================================
# GENERAL WRITER HELPERS
# ============================================================

def add_heading(document, text, level=None):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(str(text))
    run.bold = True
    run.font.size = Pt(12)
    return paragraph


def _format_percentage(value):
    return "N/A" if value is None else f"{value:.2f}%"


def add_frequency_table(document, title, data):
    add_heading(document, title)

    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"

    for cell, header in zip(
        table.rows[0].cells,
        ["Category", "Frequency", "Percentage"],
    ):
        cell.text = header

    for item in data["results"]:
        row = table.add_row().cells
        row[0].text = str(item["Category"])
        row[1].text = str(item["Frequency"])
        row[2].text = _format_percentage(item["Percentage"])

    row = table.add_row().cells
    row[0].text = "TOTAL"
    row[1].text = str(data["category_total"])
    row[2].text = _format_percentage(data["percentage_total"])

    row = table.add_row().cells
    row[0].text = "TOTAL CHECK"
    row[1].text = f"{data['category_total']} / {data['total']}"
    row[2].text = "PASS" if data["validation"] else "FAIL"

    document.add_paragraph()


# ============================================================
# CITIZEN'S CHARTER
# ============================================================

def add_cc_table(document, title, data):
    add_heading(document, title)

    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for cell, header in zip(
        table.rows[0].cells,
        ["Rating", "Frequency", "Percentage"],
    ):
        cell.text = header

    for item in data["results"]:
        row = table.add_row().cells
        row[0].text = str(item["Rating"])
        row[1].text = str(item["Frequency"])
        row[2].text = _format_percentage(item["Percentage"])

    row = table.add_row().cells
    row[0].text = "TOTAL RESPONSES"
    row[1].text = str(data["total_responses"])
    row[2].text = (
        "100.00%" if data["total_responses"] > 0 else "N/A"
    )

    row = table.add_row().cells
    row[0].text = "VALID RESPONSES"
    row[1].text = str(data["valid_responses"])
    row[2].text = ""

    row = table.add_row().cells
    row[0].text = "TOTAL CHECK"
    row[1].text = (
        f"{data['calculated_total']} / "
        f"{data['total_responses']}"
    )
    row[2].text = "PASS" if data["validation"] else "FAIL"

    document.add_paragraph()


def add_cc_item_table(document, item_name, data):
    document.add_heading(item_name, level=3)

    distribution = data["distribution"]

    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"

    for cell, header in zip(
        table.rows[0].cells,
        ["Rating", "Frequency", "Percentage"],
    ):
        cell.text = header

    for item in distribution["results"]:
        row = table.add_row().cells
        row[0].text = str(item["Rating"])
        row[1].text = str(item["Frequency"])
        row[2].text = _format_percentage(item["Percentage"])

    row = table.add_row().cells
    row[0].text = "TOTAL RESPONSES"
    row[1].text = str(distribution["total_responses"])
    row[2].text = (
        "100.00%" if distribution["total_responses"] > 0 else "N/A"
    )

    row = table.add_row().cells
    row[0].text = "VALID RESPONSES"
    row[1].text = str(distribution["valid_responses"])
    row[2].text = ""

    row = table.add_row().cells
    row[0].text = "TOTAL CHECK"
    row[1].text = (
        f"{distribution['calculated_total']} / "
        f"{distribution['total_responses']}"
    )
    row[2].text = (
        "PASS" if distribution["validation"] else "FAIL"
    )

    document.add_paragraph()


def add_cc_total_table(document, cc_name, data):
    document.add_heading(
        f"{cc_name} Total Score",
        level=3,
    )

    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"

    for cell, header in zip(
        table.rows[0].cells,
        ["Measure", "Result", "Percentage"],
    ):
        cell.text = header

    row = table.add_row().cells
    row[0].text = "Total CC Score"
    row[1].text = (
        "N/A"
        if data.get("score") is None
        else f"{data['score']:.2f}"
    )
    row[2].text = _format_percentage(data.get("percentage"))

    row = table.add_row().cells
    row[0].text = "Total Percentage"
    row[1].text = _format_percentage(data.get("percentage"))
    row[2].text = "Sum of configured item percentages"

    row = table.add_row().cells
    row[0].text = "Interpretation"
    row[1].text = str(data.get("interpretation", "N/A"))
    row[2].text = ""

    if data.get("items_used"):
        row = table.add_row().cells
        row[0].text = "Items Used"
        row[1].text = " + ".join(map(str, data["items_used"]))
        row[2].text = ""

    missing = data.get("missing_items")

    if missing:
        row = table.add_row().cells
        row[0].text = "Missing Items"
        row[1].text = ", ".join(map(str, missing))
        row[2].text = "TOTAL NOT CALCULATED"

    document.add_paragraph()


# ============================================================
# SQD
# ============================================================

def add_sqd_table(document, sqd_name, data):
    document.add_heading(sqd_name, level=2)

    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"

    for cell, header in zip(
        table.rows[0].cells,
        ["Rating", "Frequency", "Percentage"],
    ):
        cell.text = header

    ratings = [
        ("5 - Very Satisfactory", data["5 - VS"]),
        ("4 - Satisfactory", data["4 - S"]),
        (
            "3 - Neither Satisfactory Nor Dissatisfactory",
            data["3 - Neither"],
        ),
        ("2 - Dissatisfactory", data["2 - DS"]),
        ("1 - Very Dissatisfactory", data["1 - VDS"]),
        ("N/A", data["N/A"]),
    ]

    total_respondents = data["Total Respondents"]

    for rating_name, frequency in ratings:
        row = table.add_row().cells

        row[0].text = rating_name
        row[1].text = str(frequency)

        row[2].text = (
            f"{(frequency / total_respondents) * 100:.2f}%"
            if total_respondents > 0
            else "N/A"
        )

    row = table.add_row().cells
    row[0].text = "TOTAL RESPONDENTS"
    row[1].text = str(total_respondents)
    row[2].text = (
        "100.00%" if total_respondents > 0 else "N/A"
    )

    row = table.add_row().cells
    row[0].text = "VALID RESPONSES"
    row[1].text = str(data["Valid Responses"])
    row[2].text = ""

    row = table.add_row().cells
    row[0].text = "AVERAGE PERCENTAGE"
    row[1].text = _format_percentage(
        data["Average Percentage"]
    )
    row[2].text = _format_percentage(
        data["Average Percentage"]
    )

    row = table.add_row().cells
    row[0].text = "INTERPRETATION"
    row[1].text = str(data["Interpretation"])
    row[2].text = ""

    row = table.add_row().cells
    row[0].text = "TOTAL CHECK"
    row[1].text = (
        f"{data['Calculated Total']} / "
        f"{total_respondents}"
    )
    row[2].text = (
        "PASS"
        if data["Validation"]
        else "FAIL"
    )

    document.add_paragraph()


# ============================================================
# REPORT
# ============================================================

def create_word_report(
    report_data,
    output_path,
    office_name="",
    reporting_period="",
):
    """
    Create one Word report for one client group.

    Parameters:
        report_data:
            Calculated report data for either Internal or External.

        output_path:
            Destination .docx file.

        office_name:
            Office name to display in the report.

        reporting_period:
            Monthly or semestral reporting period.
    """

    document = Document()

    # --------------------------------------------------------
    # REPORT HEADER
    # --------------------------------------------------------

    document.add_heading(
        "CLIENT SATISFACTION MEASUREMENT (CSM) REPORT",
        level=1,
    )

    document.add_paragraph(
        "Caraga State University"
    )

    if office_name:
        document.add_paragraph(
            f"Office: {office_name}"
        )

    if reporting_period:
        document.add_paragraph(
            f"Reporting Period: {reporting_period}"
        )

    document.add_paragraph(
        f"Total Respondents: "
        f"{report_data.get('total_respondents', 0)}"
    )

    # --------------------------------------------------------
    # DEMOGRAPHICS
    # --------------------------------------------------------

    for key, title in [
        ("name_of_service", "Name of Service"),
        ("client_type", "Client Type"),
        ("sex", "Sex"),
        ("age", "Age"),
        ("region", "Region of Residence"),
    ]:

        if key in report_data:
            add_frequency_table(
                document,
                title,
                report_data[key],
            )

    # --------------------------------------------------------
    # CITIZEN'S CHARTER
    # --------------------------------------------------------

    if "citizen_charter" in report_data:

        document.add_heading(
            "Citizen's Charter",
            level=1,
        )

        for cc_name, cc in report_data[
            "citizen_charter"
        ].items():

            document.add_heading(
                cc_name,
                level=2,
            )

            # ------------------------------------------------
            # LEGACY CC DISTRIBUTION
            # ------------------------------------------------

            legacy = cc.get(
                "legacy_distribution"
            )

            if legacy is not None:
                add_cc_table(
                    document,
                    f"{cc_name} Distribution",
                    legacy,
                )

            # ------------------------------------------------
            # ITEM-LEVEL CC TABLES
            # ------------------------------------------------

            for item_name, item_data in cc.get(
                "items",
                {}
            ).items():

                add_cc_item_table(
                    document,
                    item_name,
                    item_data,
                )

            # ------------------------------------------------
            # TOTAL CC SCORE
            # ------------------------------------------------

            total_score = cc.get(
                "total_score"
            )

            if isinstance(total_score, dict):

                # Show total if it is calculable.
                #
                # If required items are missing,
                # show the N/A/missing-items result
                # transparently instead of fabricating
                # a value.

                if (
                    total_score.get("score") is not None
                    or total_score.get("missing_items")
                ):

                    add_cc_total_table(
                        document,
                        cc_name,
                        total_score,
                    )

    # --------------------------------------------------------
    # SERVICE QUALITY DIMENSIONS
    # --------------------------------------------------------

    if "sqd" in report_data:

        document.add_heading(
            "Service Quality Dimensions",
            level=1,
        )

        for sqd_name, sqd_data in report_data[
            "sqd"
        ].items():

            add_sqd_table(
                document,
                sqd_name,
                sqd_data,
            )

    # --------------------------------------------------------
    # SAVE DOCUMENT
    # --------------------------------------------------------

    document.save(output_path)